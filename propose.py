# -*- coding: utf-8 -*-
"""propose.py - בונה הצעת מערכת בשם, בצד, בלי לגעת במערכת המפורסמת.

   שימוש:  python propose.py <שם> [שניות לפותר]
   ההצעה יושבת ב-proposals/<שם>/proposal.json:
   {
     "title":  "הצעה: ...",            כותרת הלוח
     "banner": "הצעה - לא מפורסם",     הסרט ליד הכותרת
     "base":   "published" | "<שם הצעה אחרת>",   ממה יוצאים (יציבות + סימוני שינוי)
     "env":    {"STAB":"80", ...},      משתני סביבה למנוע (כפתורי ניסוי)
     "rules_off": ["track_day"],        חוקי מדיניות לכיבוי (rules.py)
     "freeze_elem": true, "freeze_hat": true,   להקפיא צד שלם על הבסיס
     "pins":   [{"class":..,"day":..,"hour":..,"teacher":..,"subject":..,"value":1}],
     "overrides": {"hdata": {...}, "data2": {...}, "data": {...}},   עקיפות נתונים (overrides.py)
     "extra":  "extra.py"               אילוצים נוספים בפייתון (רץ בתוך המנוע)
   }
   הפלט ב-proposals/<שם>/out/: viewer.html (עם סימוני היה ← נהיה), impact.txt
   (היקף ההשפעה: תאים, כיתות, מורים), sol_*.json, אקסל עם גיליון "שינויים".
   הבנייה ב-proposals/<שם>/build/ - עותק של הריפו, נזרק בכל ריצה.
   שום דבר כאן לא נוגע ב-master, ב-baseline או ב-viewer.html של הריפו."""
import io, json, os, shutil, subprocess, sys, collections
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = os.path.dirname(os.path.abspath(__file__))

def sh(args, cwd, env=None, label=""):
    r = subprocess.run(args, cwd=cwd, env=env, capture_output=True, text=True, encoding="utf-8", errors="replace")
    out = (r.stdout or "") + (r.stderr or "")
    tail = [l for l in (r.stdout or "").strip().split("\n") if l.strip()][-2:]
    if label: print("  " + label + ": " + " | ".join(tail))
    return r.returncode, out

def main():
    if len(sys.argv) < 2: print(__doc__); sys.exit(2)
    name = sys.argv[1]; TL = sys.argv[2] if len(sys.argv) > 2 else "200"
    pdir = os.path.join(ROOT, "proposals", name)
    cfg = json.load(io.open(os.path.join(pdir, "proposal.json"), encoding="utf-8"))
    bdir = os.path.join(pdir, "build"); odir = os.path.join(pdir, "out")
    shutil.rmtree(bdir, ignore_errors=True); os.makedirs(bdir); os.makedirs(odir, exist_ok=True)

    # 1. עותק נקי של הריפו (קבצים במעקב, בלי PDF) + קלטים מקומיים שאינם בגיט
    files = subprocess.run(["git", "ls-files", "-z", "--cached", "--others", "--exclude-standard"], cwd=ROOT, capture_output=True).stdout.decode("utf-8").split("\0")
    for f in files:
        if not f or f.endswith(".pdf") or f.startswith("proposals/"): continue
        dst = os.path.join(bdir, f); os.makedirs(os.path.dirname(dst), exist_ok=True); shutil.copy2(os.path.join(ROOT, f), dst)
    for f in ("ids_local.json", "app_data"):
        src = os.path.join(ROOT, f)
        if os.path.isdir(src): shutil.copytree(src, os.path.join(bdir, f))
        elif os.path.exists(src): shutil.copy2(src, os.path.join(bdir, f))

    # 2. בסיס: המפורסם (baseline_* של הריפו) או הצעה אחרת (ה-out שלה)
    base = cfg.get("base", "published")
    base_viewer = os.path.join(ROOT, "viewer.html"); mid_viewer = None
    if base != "published":
        bo = os.path.join(ROOT, "proposals", base, "out")
        for a, b in (("sol_J.json", "baseline_J.json"), ("sol_hat.json", "baseline_hat.json"), ("tln_map.json", "baseline_tln.json")):
            if os.path.exists(os.path.join(bo, a)): shutil.copy2(os.path.join(bo, a), os.path.join(bdir, b))
        mid_viewer = os.path.join(bo, "viewer.html")
        print(f"בסיס: הצעה '{base}'")

    # 3. סביבה למנוע
    env = dict(os.environ); env.update({k: str(v) for k, v in cfg.get("env", {}).items()}); env["TL"] = TL; env["NODIAG"] = "1"
    if cfg.get("rules_off"): env["RULES_OFF"] = ",".join(cfg["rules_off"])
    # הקפאה: true = על הבסיס; "<קובץ>" = על פתרון שמור בתיקיית ההצעה (שחזור מדויק של הצעה שכבר הוצגה)
    for key, envk, tmp in (("freeze_elem", "FREEZEJ", "_freeze_J.json"), ("freeze_hat", "FREEZEH", "_freeze_H.json")):
        fz = cfg.get(key)
        if fz is True: env[envk] = "1"
        elif fz: shutil.copy2(os.path.join(pdir, fz), os.path.join(bdir, tmp)); env[envk] = tmp
    if cfg.get("pins"):
        json.dump(cfg["pins"], io.open(os.path.join(bdir, "_pins.json"), "w", encoding="utf-8"), ensure_ascii=False); env["PINS"] = "_pins.json"
    if cfg.get("overrides"):
        json.dump(cfg["overrides"], io.open(os.path.join(bdir, "_overrides.json"), "w", encoding="utf-8"), ensure_ascii=False); env["OVERRIDES"] = "_overrides.json"
    if cfg.get("extra"):
        shutil.copy2(os.path.join(pdir, cfg["extra"]), os.path.join(bdir, "_extra.py")); env["EXTRA"] = "_extra.py"

    # 4. הצינור (בלי PDF ובלי API - זו הצעה, לא פרסום)
    print(f"בונה הצעה '{name}' (מגבלת זמן {TL} שניות)...")
    rc, out = sh([sys.executable, "engine.py"], bdir, env, "פותר")
    io.open(os.path.join(odir, "engine.log"), "w", encoding="utf-8").write(out)
    if rc != 0 and "INFEASIBLE" not in out:
        print("\n!!! המנוע קרס (לא בעיית אילוצים). הלוג ב-out/engine.log"); print(out[-2000:]); sys.exit(1)
    if rc != 0 or "INFEASIBLE" in out or "MODEL_INVALID" in out:
        print("\n!!! הפותר לא מצא פתרון להצעה. הלוג ב-out/engine.log")
        if "קיבולת" in out: print(out[out.find("קיבולת"):][:1500])
        print("מריץ אבחון (הרפיה אחת בכל פעם)...")
        rc2, out2 = sh([sys.executable, "diagnose.py", "45"], bdir, env); print(out2[-3000:])
        sys.exit(1)
    for script, label in (("fill2.py", "ממלא חורים"), ("checks.py", "בדיקות"), ("outGAPS.py", "אקסל")):
        rc, o = sh([sys.executable, script], bdir, env, label)
        if rc != 0: print(o[-1500:]); sys.exit(1)
    env["PROP_BASE"] = base_viewer; env["PROP_TITLE"] = cfg.get("title", name); env["PROP_BANNER"] = cfg.get("banner", "הצעה - לא מפורסם")
    if mid_viewer: env["PROP_MID"] = mid_viewer
    rc, o = sh([sys.executable, "make_viewer.py"], bdir, env, "לוח")
    if rc != 0: print(o[-1500:]); sys.exit(1)

    # 5. היקף ההשפעה: מול המפורסם תמיד; ומול הבסיס אם הוא הצעה
    sys.path.insert(0, ROOT)
    from data2 import CLASSES, SLOTS, DAY_NAMES
    from hdata import HCLASSES
    SJ = json.load(io.open(os.path.join(bdir, "sol_J.json"), encoding="utf-8")); SH = json.load(io.open(os.path.join(bdir, "sol_hat.json"), encoding="utf-8"))
    HSL = sorted({tuple(map(int, k.split(","))) for c in SH for k in SH[c]})
    def diff(cur, prev, classes, slots):
        return [(c, d, h, prev.get(c, {}).get(f"{d},{h}") or "", cur.get(c, {}).get(f"{d},{h}") or "")
                for c in classes for (d, h) in slots if (prev.get(c, {}).get(f"{d},{h}") or "") != (cur.get(c, {}).get(f"{d},{h}") or "")]
    def tname(v): return v.split(" – ")[-1] if " – " in v else v
    def report(title, bJ, bH):
        dj = diff(SJ, bJ, CLASSES, SLOTS); dh = diff(SH, bH, HCLASSES, HSL)
        L = [f"== {title} ==", f"יסודי: {len(dj)} תאים | חטיבה: {len(dh)} תאים"]
        byc = collections.Counter(c for c, *_ in dj + dh)
        L.append("לפי כיתה: " + (", ".join(f"{c} {n}" for c, n in sorted(byc.items())) or "—"))
        ts = collections.Counter()
        for c, d, h, a, b in dj + dh:
            for v in (a, b):
                for t in tname(v).split(" + "):
                    if t: ts[t] += 1
        L.append("מורים מעורבים: " + (", ".join(f"{t} ({n})" for t, n in ts.most_common()) or "—"))
        for c, d, h, a, b in sorted(dh) + sorted(dj):
            L.append(f"  {c} | {DAY_NAMES[d]} ש{h}: {a or '—'}  ←  {b or '—'}")
        # שעות לפי יום למורים המעורבים: לפני ← אחרי
        def perday(sol_j, sol_h, t):
            n = [0] * 6
            for c in sol_j:
                for k, v in sol_j[c].items():
                    if v == t: n[int(k.split(",")[0])] += 1
            for c in sol_h:
                for k, v in sol_h[c].items():
                    if t in tname(v or "").split(" + "): n[int(k.split(",")[0])] += 1
            return n
        if ts:
            L.append("שעות לפי יום (א ב ג ד ה ו): לפני ← אחרי")
            for t in sorted(ts):
                if t in ("חסר מורה", "מגמות", 'תל"ן'): continue
                p, q = perday(bJ, bH, t), perday(SJ, SH, t)
                L.append(f"  {t:8s} {' '.join(map(str, p))} (סה\"כ {sum(p)})  ←  {' '.join(map(str, q))} (סה\"כ {sum(q)})")
        return L, dj + dh
    BJ = json.load(io.open(os.path.join(ROOT, "baseline_J.json"), encoding="utf-8")); BH = json.load(io.open(os.path.join(ROOT, "baseline_hat.json"), encoding="utf-8"))
    lines, cells = report("מול המערכת המפורסמת", BJ, BH)
    if base != "published":
        bJ = json.load(io.open(os.path.join(bdir, "baseline_J.json"), encoding="utf-8")); bH = json.load(io.open(os.path.join(bdir, "baseline_hat.json"), encoding="utf-8"))
        l2, _ = report(f"מול הבסיס (הצעה '{base}')", bJ, bH); lines += [""] + l2
    text = "\n".join(lines)
    io.open(os.path.join(odir, "impact.txt"), "w", encoding="utf-8").write(text)
    print("\n" + "\n".join(lines[:4]) + f"\n  (הרשימה המלאה ב-out/impact.txt)")

    # 6. אקסל: המערכות של ההצעה + גיליון "שינויים"
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = load_workbook(os.path.join(bdir, "מערכות שעות.xlsx"))
        ws = wb.create_sheet("שינויים", 0); ws.sheet_view.rightToLeft = True
        for col, w in zip("ABCDE", (14, 10, 6, 30, 30)): ws.column_dimensions[col].width = w
        ws.append([cfg.get("title", name)]); ws["A1"].font = Font(bold=True, size=13)
        for l in lines[1:4]: ws.append([l])
        ws.append([]); ws.append(["כיתה", "יום", "שעה", "היה", "נהיה"])
        for c in ws[ws.max_row]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="EEE8DC")
        for c, d, h, a, b in sorted(cells, key=lambda z: (z[0] in CLASSES, z[0], z[1], z[2])):
            ws.append([c, DAY_NAMES[d], h, a or "—", b or "—"])
        wb.save(os.path.join(odir, f"מערכות שעות - {name}.xlsx")); print(f"  אקסל: out/מערכות שעות - {name}.xlsx")
    except Exception as e:
        print("  אקסל: דילוג (" + str(e) + ")")

    for f in ("viewer.html", "sol_J.json", "sol_hat.json", "tln_map.json", "fills.json"):
        if os.path.exists(os.path.join(bdir, f)): shutil.copy2(os.path.join(bdir, f), os.path.join(odir, f))
    print(f"\nהצעה '{name}' מוכנה: proposals/{name}/out/viewer.html  (לא מפורסם - רק אחרי אישור)")

if __name__ == "__main__": main()
