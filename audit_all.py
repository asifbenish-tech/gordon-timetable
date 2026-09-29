# -*- coding: utf-8 -*-
"""audit_all.py - בדיקה מקיפה של לוח (מפורסם או הצעה) + האקסל שלו.
   python audit_all.py <viewer.html> [<xlsx>] [<build dir לנתונים/עקיפות>]
   בודק: מורה בשני מקומות באותה שעה, שיעור מול סדירות, שעות לפי מקצוע מול
   תוכנית (חטיבה), שעות מורה-כיתה מול מכסה (יסודי), ניצול מול מכסה, התאמה
   תא-בתא בין האקסל ללוח, ומורה שמופיע פעמיים בגיליון המורים."""
import io, json, re, sys, collections, os
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
V = os.path.abspath(sys.argv[1]); X = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else None
if len(sys.argv) > 3:
    sys.path.insert(0, os.path.abspath(sys.argv[3])); os.chdir(sys.argv[3])
from data import QUOTA as EQ
from hdata import NEED, OVR, GRADE, JOINT
D = json.loads(re.search(r"^const DATA=(\{.*\});?$", io.open(V, encoding="utf-8").read(), re.M).group(1))
err, warn = [], []
DN = D["days"]
jointkeys = {(j["teacher"], j["day"], j["hour"]) for j in JOINT}

# 1. מורה בשני מקומות באותה שעה / שיעור בזמן סדירות
for t, ev in D["teachers"].items():
    if t in ("מגמות", "חסר מורה", "שרית + חסן", "שכבת ט יחד"): continue
    by = collections.defaultdict(list)
    for side, d, h, lbl in ev: by[(d, h)].append((side, lbl))
    for (d, h), items in sorted(by.items()):
        real = [x for x in items if x[0] != "סדירות"]
        sed = [x for x in items if x[0] == "סדירות"]
        subj = {x[1].split(" · ")[1] if " · " in x[1] else "" for x in real}
        legit = len(real) > 1 and (subj <= {"חינוך גופני", "שירה בציבור"} or (t, d, h) in jointkeys)
        if len(real) > 1 and not legit:
            err.append(f"{t}: {DN[d]} ש{h} בכמה מקומות: " + " | ".join(x[1] for x in real))
        if real and sed:
            warn.append(f"{t}: {DN[d]} ש{h} שיעור ({real[0][1]}) בזמן סדירות ({sed[0][1]})")

# 2. חטיבה: שעות לפי מקצוע מול תוכנית הלימודים
for c, info in D["jun"].items():
    cnt = collections.Counter(v.get("t") for v in info["cells"].values() if v.get("t") and v.get("k") != "off")
    g = GRADE[c]
    for sj, per in NEED.items():
        want = OVR.get((c, sj), per[g]); got = cnt.get(sj, 0)
        if want and got != want: warn.append(f"{c}: {sj} {got} שעות בלוח, בתוכנית {want}")

# 3. יסודי: שעות מורה-כיתה מול המכסה
act = collections.Counter()
for c, info in D["elem"].items():
    for k, v in info["cells"].items():
        if v.get("t"): act[(v["t"], c)] += 1
for t, cls in EQ.items():
    for c, q in cls.items():
        if act[(t, c)] != q and t != 'תל"ן':
            warn.append(f"יסודי: {t} ב{c} {act[(t, c)]} שעות בלוח, מכסה {q}")

# 4. ניצול מול מכסה
for u in D["util"]:
    if u["left"] < 0: err.append(f"{u['t']}: משובץ {u['tot']} מעל המכסה {u['q']}")

# 5. אקסל מול הלוח
if X:
    from openpyxl import load_workbook
    wb = load_workbook(X); nx = 0
    for side in ("elem", "jun"):
        for c, info in D[side].items():
            if c not in wb.sheetnames: err.append(f"אקסל: אין גיליון לכיתה {c}"); continue
            rows = {r[0].value: r for r in wb[c].iter_rows(min_row=3) if isinstance(r[0].value, int)}
            for k, v in info["cells"].items():
                d, h = map(int, k.split(",")); t = v.get("t") or ""; s = v.get("s") or ""
                if not t: continue
                if h not in rows: err.append(f"אקסל {c}: חסרה שורת שעה {h}"); continue
                xv = str(rows[h][1 + d].value or "").replace(" – ", " · ")
                if t.replace(" – ", " · ") not in xv and not (s and s in xv):
                    nx += 1
                    if nx <= 20: err.append(f"אקסל {c} {DN[d]} ש{h}: בלוח '{t} / {s}', באקסל '{xv}'")
    if nx > 20: err.append(f"... ועוד {nx-20} אי-התאמות אקסל/לוח")
    names = [r[0] for r in wb["מערכות מורים"].iter_rows(values_only=True)
             if isinstance(r[0], str) and r[1] and ("משובץ" in str(r[1]) or "שעות בשבוע" in str(r[1]))]
    dup = [n for n, k in collections.Counter(names).items() if k > 1]
    if dup: err.append(f"אקסל: מורים שמופיעים פעמיים בגיליון המורים: {dup}")
    print(f"גיליון מורים: {len(names)} מורים")

print(f"== {os.path.basename(V)}{' + ' + os.path.basename(X) if X else ''}")
print(f"שגיאות: {len(err)}"); [print("  ✗", e) for e in err]
print(f"לבדיקה: {len(warn)}"); [print("  ⚠", w) for w in warn]
sys.exit(1 if err else 0)
