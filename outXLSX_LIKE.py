# -*- coding: utf-8 -*-
"""outXLSX_LIKE.py - מערכות השעות באקסל, בעיצוב הקובץ של בית החינוך.
   כותרת Arial 14 מודגשת, שורת ימים ועמודת שעות בכחול (28486B) עם טקסט לבן,
   תאים Arial 9 ממורכזים עם גלישת שורה, גיליון מימין לשמאל.
   הרצה: python outXLSX_LIKE.py  ->  'מערכות שעות - גורדון.xlsx'"""
import io, json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# לוח הצלצולים (עדכון המנהל 07.09.2026). מקור אחד לכל הגיליונות.
BELLS = [(1, "8:00-8:45"), (2, "8:45-9:30"), (3, "9:40-10:25"), (4, "10:25-11:10"),
         (5, "11:30-12:15"), (6, "12:35-13:20"), (7, "13:20-14:05")]
BREAKS = [(2, "הפסקה", "10 דק'"), (4, "הפסקה ארוכה - משמרת ראשונה", "20 דק'"),
          (5, "הפסקה ארוכה - משמרת שנייה", "20 דק'")]
TIME = dict(BELLS)
DAYS = ["ראשון", "שני", "שלישי", "רביעי", "חמישי", "שישי"]

NAVY  = PatternFill("solid", fgColor="FF28486B")
WHITE = PatternFill("solid", fgColor="FFFFFFFF")
GREY  = PatternFill("solid", fgColor="FFF2F2F2")
HEAD  = Font(name="Arial", size=9, bold=True, color="FFFFFFFF")
BODY  = Font(name="Arial", size=9, color="FF000000")
TITLE = Font(name="Arial", size=14, bold=True, color="FF000000")
_thin = Side(style="thin")
BOX   = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
MID   = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="bottom")

C0 = 3                                    # הבלוק מתחיל בעמודה C, כמו בקובץ המקורי


def _sheet(wb, name):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    for i in range(1, C0 + 8):
        ws.column_dimensions[CL(i)].width = 13
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 3
    return ws


def block(ws, row, title, hours, cell_at):
    """בלוק מערכת אחד. cell_at(d,h) -> טקסט התא (או "")."""
    ws.cell(row, C0, title).font = TITLE
    ws.cell(row, C0).alignment = RIGHT
    ws.merge_cells(start_row=row, start_column=C0, end_row=row, end_column=C0 + 2)
    row += 1
    for j, lbl in enumerate(["שעה / יום"] + DAYS):
        c = ws.cell(row, C0 + j, lbl)
        c.font, c.fill, c.border, c.alignment = HEAD, NAVY, BOX, MID
    row += 1
    c = ws.cell(row, C0, "בוקר טוב")
    c.font, c.fill, c.border, c.alignment = HEAD, NAVY, BOX, MID
    for j in range(1, 7):
        c = ws.cell(row, C0 + j)
        c.fill, c.border, c.alignment = WHITE, BOX, MID
    row += 1
    for h in hours:
        c = ws.cell(row, C0, f"{h}\n {TIME[h]}")
        c.font, c.fill, c.border, c.alignment = HEAD, NAVY, BOX, MID
        for d in range(6):
            c = ws.cell(row, C0 + 1 + d, cell_at(d, h) or None)
            c.font, c.fill, c.border, c.alignment = BODY, WHITE, BOX, MID
        # בלי גובה קבוע - אקסל מתאים את הגובה לטקסט, כמו בקובץ המקורי
        row += 1
    return row + 1                        # שורה ריקה בין בלוקים


def bells_sheet(wb):
    ws = _sheet(wb, "לוח צלצולים")
    ws.cell(1, C0, "לוח צלצולים").font = TITLE
    ws.cell(1, C0).alignment = RIGHT
    ws.merge_cells(start_row=1, start_column=C0, end_row=1, end_column=C0 + 2)
    for j, lbl in enumerate(["שעה", "משעה", "עד שעה", "משך"]):
        c = ws.cell(2, C0 + j, lbl)
        c.font, c.fill, c.border, c.alignment = HEAD, NAVY, BOX, MID
    r = 3
    brk = {h: (n, d) for h, n, d in BREAKS}
    for h, rng in BELLS:
        a, b = rng.split("-")
        for j, v in enumerate((f"שעה {h}", a, b, "45 דק'")):
            c = ws.cell(r, C0 + j, v)
            c.font, c.fill, c.border, c.alignment = BODY, WHITE, BOX, MID
        r += 1
        if h in brk:
            name, dur = brk[h]
            c = ws.cell(r, C0, name)
            c.font, c.fill, c.border, c.alignment = BODY, GREY, BOX, MID
            ws.merge_cells(start_row=r, start_column=C0, end_row=r, end_column=C0 + 2)
            c = ws.cell(r, C0 + 3, dur)
            c.font, c.fill, c.border, c.alignment = BODY, GREY, BOX, MID
            r += 1
    ws.column_dimensions[CL(C0)].width = 26
    return ws


def main():
    s = io.open("viewer.html", encoding="utf-8").read()
    D = json.loads(re.search(r"^const DATA=(\{.*\});?$", s, re.M).group(1))

    wb = Workbook()
    wb.remove(wb.active)

    # ---------- מערכות הכיתות ----------
    ws = _sheet(wb, "מערכות כיתה")
    row = 1
    for cname, info in list(D["elem"].items()) + list(D["jun"].items()):
        cells = info["cells"]
        hrs = sorted({int(k.split(",")[1]) for k in cells})

        def at(d, h, cells=cells):
            v = cells.get(f"{d},{h}")
            if not v or not v.get("t"):
                return ""
            head, sub = v["t"], (v.get("s") or "")
            if v.get("co"):
                # בלי אמוג'י באקסל - לא כל גופן להדפסה מציג אותו
                sub = (sub + " · " if sub else "") + v["co"].replace(" \U0001F37D", "")
            return f"{head}\n {sub}" if sub else head

        row = block(ws, row, f"מערכת שעות כיתה {cname}", hrs, at)

    # ---------- מערכות המורים ----------
    ws = _sheet(wb, "מערכות מורים")
    row = 1
    for t in sorted(D["teachers"]):
        if t in ("מגמות", "חסר מורה", "שרית + חסן", "שכבת ט יחד"):
            continue
        ev = {}
        for _lvl, d, h, lbl in D["teachers"][t]:
            if _lvl == "סדירות":            # מחויבות שאינה הוראה - מסומנת כמו בלוח
                ev.setdefault((d, h), []).append("◦ " + lbl)
                continue
            # בקובץ המקורי: המקצוע בשורה הראשונה והכיתה בשנייה. היפוך רק
            # כששני חלקים ("ט אסיף · אנגלית"); ברשימות המגמות הסדר נשמר.
            parts = lbl.split(" · ")
            head, sub = (parts[1], parts[0]) if len(parts) == 2 else (parts[0], " · ".join(parts[1:]))
            ev.setdefault((d, h), []).append(f"{head}\n {sub}" if sub else head)
        full = (D.get("full_names") or {}).get(t, t)
        row = block(ws, row, f"מערכת שעות למורה {full}",
                    [h for h, _ in BELLS],
                    lambda d, h, ev=ev: " / ".join(ev.get((d, h), [])))

    bells_sheet(wb)
    wb._sheets = [wb["לוח צלצולים"], wb["מערכות כיתה"], wb["מערכות מורים"]]
    out = "מערכות שעות - גורדון.xlsx"
    wb.save(out)
    print(f"{out} נוצר: {len(wb.sheetnames)} גיליונות")


if __name__ == "__main__":
    main()
